"""
Wish Karo listing generator.

    python generate.py  Master.xlsx  [output_folder]

Reads the filled master sheet and writes, into output_folder (default: ./output):
    Meesho_<date>.xlsx   Flipkart_<date>.xls   Myntra_<date>.xlsx   Amazon_<date>.xlsm
    Listing_Report.xlsx  -> everything that needs a human look before upload
Each file is the marketplace's own template, filled in, with dropdowns intact.
"""
import re
import datetime as dt
import json
import os
import sys
import warnings
from collections import OrderedDict, defaultdict

import openpyxl
import xlrd

# openpyxl only READS the master and templates here (output files are written by writers.py,
# which keeps every dropdown), so its "extension not supported" notice is harmless noise.
warnings.filterwarnings("ignore", message="Data Validation extension is not supported", category=UserWarning)

import master_spec as S
import vocab as V
import writers

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.path.join(HERE, "templates")
ALLOWED = json.load(open(os.path.join(HERE, "config", "allowed_values.json"), encoding="utf-8"))

TEMPLATES = {
    "meesho":   dict(file="meesho.xlsx",  sheet="Frocks---Dresses-Fill this", header_row=3, data_row=5, ext="xlsx"),
    "flipkart": dict(file="flipkart.xls", sheet="kids_dress",                 header_row=1, data_row=5, ext="xls"),
    "myntra":   dict(file="myntra.xlsx",  sheet="Dresses",                    header_row=3, data_row=4, ext="xlsx"),
    # Amazon's newer template: row 4 = labels, row 5 = attribute keys, row 6 = Amazon's example, data from row 7
    "amazon":   dict(file="amazon.xlsm",  sheet="Template",                   header_row=5, data_row=7, ext="xlsm"),
}
CH_NAME = {"master": "Master sheet", "meesho": "Meesho", "flipkart": "Flipkart", "myntra": "Myntra", "amazon": "Amazon"}


# ======================================================== issue collector ===
class Report:
    def __init__(self):
        self.rows = []

    def add(self, level, channel, style, sku, field, msg):
        self.rows.append((level, CH_NAME.get(channel, channel), style or "", sku or "", field, msg))

    def err(self, *a): self.add("ERROR", *a)
    def warn(self, *a): self.add("CHECK", *a)
    def info(self, *a): self.add("INFO", *a)


# ============================================================ read master ===
def _clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


NUMERIC_KEYS = {"meesho_price", "meesho_mrp", "flipkart_price", "flipkart_mrp", "myntra_price", "myntra_mrp",
                "amazon_price", "amazon_mrp", "stock", "chest", "front_length", "bust", "waist", "hips",
                "shoulder", "pkg_l", "pkg_w", "pkg_h", "weight_g"}
IMG_KEYS = {"img_front", "img_back", "img_side", "img_detail", "img_look", "img6", "img7", "img8",
            "meesho_img1", "meesho_img2", "meesho_img3", "meesho_img4"}
NO_SPACE_KEYS = {"style_code", "parent_code", "sku", "hsn"}


def _norm_size(v):
    """'6-7 years', '6 - 7 Years', '6-7yrs', '3-6m' -> '6-7 Years' / '3-6 Months'"""
    t = re.sub(r"\s+", " ", str(v)).strip().lower()
    m = re.match(r"^(\d+)\s*-\s*(\d+)\s*(y|yr|yrs|year|years|m|mo|mos|month|months)$", t)
    if not m:
        return v
    unit = "Years" if m.group(3).startswith("y") else "Months"
    return f"{m.group(1)}-{m.group(2)} {unit}"


def auto_fix(key, v, col):
    """Clean typing/spacing mistakes. Returns the fixed value."""
    if v is None:
        return None
    if isinstance(v, str):
        v = v.replace("\u00a0", " ").replace("\u200b", "")
        if key in IMG_KEYS:
            v = re.sub(r"\s+", "%20", v.strip())             # a space inside a link is part of the file name
        elif key in NO_SPACE_KEYS:
            v = re.sub(r"\s+", "", v)
        elif key == "description":
            v = "\n".join(re.sub(r"[ \t]+", " ", line).strip() for line in v.splitlines()).strip()
            v = re.sub(r"\n{3,}", "\n\n", v)
        else:
            v = re.sub(r"\s+", " ", v).strip()
            v = re.sub(r"\s+([,.!?])", r"\1", v)          # "dress ," -> "dress,"
        if v == "":
            return None
    if key in NUMERIC_KEYS and isinstance(v, str):
        kilos = key == "weight_g" and re.search(r"(?i)\d\s*kgs?$", v)    # "0.3 kg" in the grams column
        t = re.sub(r"(?i)rs\.?|inr|₹|,|cm|kg|g|inches|inch|in|\"", "", v).strip()
        try:
            f = round(float(t) * 1000, 1) if kilos else float(t)
            v = int(f) if f.is_integer() else f
        except ValueError:
            pass
    if key == "size" and v is not None:
        v = _norm_size(v)
    if col and col.get("list") and isinstance(v, (str, int, float)):
        opts = S.LISTS.get(col["list"]) or []
        if v not in opts:
            low = {str(o).strip().lower(): o for o in opts}
            hit = low.get(str(v).strip().lower())
            if hit is None and col["list"] == "sizes":
                hit = low.get(str(_norm_size(v)).lower())
            if hit is not None:
                v = hit
    return v


def read_table(ws, groups, fixes=None):
    spec = {c["header"]: c for _, cols in groups for c in cols}
    by_key = {c["key"]: c for _, cols in groups for c in cols}
    headers = {}
    for ci in range(1, ws.max_column + 1):
        h = ws.cell(2, ci).value
        if h:
            h = str(h).replace(" *", "").strip()
            if h in spec:
                headers.setdefault(h, []).append(ci)
    # the Sizes sheet repeats price headers for overrides -> take the right one
    keys = {}
    for _, cols in groups:
        for c in cols:
            idx = headers.get(c["header"], [])
            if idx:
                keys[c["key"]] = idx.pop(0)
    out = []
    for r in range(4, ws.max_row + 1):
        row = {}
        for k, ci in keys.items():
            raw = _clean(ws.cell(r, ci).value)
            fixed = _clean(auto_fix(k, raw, by_key.get(k)))
            if fixes is not None and raw is not None and fixed != raw:
                fixes.append((ws.title, r, by_key[k]["header"], raw, fixed))
            row[k] = fixed
        if not any(v is not None for v in row.values()):
            continue
        row["_row"] = r
        out.append(row)
    return out


def read_settings(ws):
    per_ch, general, size_tbl = {}, {}, {}
    mode = None
    for r in range(1, ws.max_row + 1):
        a = _clean(ws.cell(r, 1).value)
        if a in ("Setting",) and _clean(ws.cell(r, 2).value) == "Meesho":
            mode = "ch"; continue
        if a == "GENERAL":
            mode = "gen"; continue
        if a == "SIZE TABLE":
            mode = "size"; continue
        if a in (None, "Setting", "Size"):
            continue
        if mode == "ch":
            per_ch[a] = {ch: _clean(ws.cell(r, 2 + i).value) for i, ch in enumerate(["meesho", "flipkart", "myntra", "amazon"])}
        elif mode == "gen":
            general[a] = _clean(ws.cell(r, 2).value)
        elif mode == "size":
            size_tbl[a] = _clean(ws.cell(r, 2).value)
    return per_ch, general, size_tbl


def load_master(path, rep):
    wb = openpyxl.load_workbook(path, data_only=True)
    for name in ("Products", "Sizes", "Settings"):
        if name not in wb.sheetnames:
            raise SystemExit(f"Master file is missing the '{name}' sheet")
    fixes = []
    products = read_table(wb["Products"], S.PRODUCT_GROUPS, fixes)
    sizes = read_table(wb["Sizes"], S.SIZE_GROUPS, fixes)
    for sheet, r, header, raw, fixed in fixes:
        if str(raw).upper().startswith("EXAMPLE"):
            continue
        rep.info("master", "", "", header,
                 f"Auto-fixed {sheet} row {r}: '{str(raw)[:60]}' -> '{str(fixed)[:60]}'")
    headers = {c["key"]: c["header"] for _, cols in S.PRODUCT_GROUPS + S.SIZE_GROUPS for c in cols}
    for sheet, rows in (("Products", products), ("Sizes", sizes)):
        for row in rows:
            for k in NUMERIC_KEYS:
                if isinstance(row.get(k), str) and not str(row.get("style_code") or "").upper().startswith("EXAMPLE"):
                    rep.err("master", row.get("style_code"), "", headers[k],
                            f"{sheet} row {row['_row']}: '{row[k]}' is not a number - left blank")
                    row[k] = None
    per_ch, general, size_tbl = read_settings(wb["Settings"])

    defaults = {c["key"]: c["default"] for _, cols in S.PRODUCT_GROUPS for c in cols if c["default"] is not None}
    prods = OrderedDict()
    for p in products:
        sc = p.get("style_code")
        if not sc or str(sc).upper().startswith("EXAMPLE"):
            continue
        sc = str(sc)
        for k, d in defaults.items():
            if p.get(k) is None:
                p[k] = d
        for k in ("hsn", "gst"):
            p[k] = str(p[k]) if p[k] is not None else None
        if sc in prods:
            # same colour typed on several rows (e.g. one per size): merge if nothing conflicts
            first = prods[sc]
            clash = [k for k, v in p.items() if not k.startswith("_") and v not in (None, "")
                     and first.get(k) not in (None, "") and str(first[k]).strip() != str(v).strip()]
            if clash:
                rep.err("master", sc, "", "Style Code",
                        f"Style Code is on rows {first['_row']} and {p['_row']} with different values in: "
                        f"{', '.join(clash[:5])}. Keep ONE row per colour.")
            else:
                for k, v in p.items():
                    if first.get(k) in (None, "") and v not in (None, ""):
                        first[k] = v
                rep.info("master", sc, "", "Style Code",
                         f"Row {p['_row']} repeats row {first['_row']} - merged. Products needs only ONE row per colour; sizes go in the Sizes sheet.")
            continue
        p["sizes"] = []
        prods[sc] = p

    seen_sku = {}
    by_lower = {k.lower(): k for k in prods}
    for z in sizes:
        sc = z.get("style_code")
        if not sc or str(sc).upper().startswith("EXAMPLE"):
            continue
        sc = str(sc)
        if sc not in prods and sc.lower() in by_lower:          # 'TST100BLU' typed for 'tst100blu'
            rep.info("master", by_lower[sc.lower()], "", "Style Code",
                     f"Auto-fixed Sizes row {z['_row']}: '{sc}' -> '{by_lower[sc.lower()]}'")
            sc = z["style_code"] = by_lower[sc.lower()]
        if sc not in prods:
            rep.err("master", sc, "", "Style Code", f"Sizes row {z['_row']}: Style Code not found in Products sheet")
            continue
        if z.get("size") not in V.SIZES:
            rep.err("master", sc, "", "Size", f"Sizes row {z['_row']}: size '{z.get('size')}' is not in the dropdown list")
            continue
        z["sku"] = str(z["sku"]) if z.get("sku") else V.sku_for(sc, z["size"])
        if z["sku"] in seen_sku:
            rep.err("master", sc, z["sku"], "SKU", f"Duplicate SKU (Sizes rows {seen_sku[z['sku']]} and {z['_row']})")
            continue
        seen_sku[z["sku"]] = z["_row"]
        prods[sc]["sizes"].append(z)

    for sc, p in prods.items():
        p["sizes"].sort(key=lambda z: V.SIZES.index(z["size"]))
        if not p["sizes"]:
            rep.err("master", sc, "", "Sizes", "Product has no rows in the Sizes sheet - skipped everywhere")
        for c in [c for _, cols in S.PRODUCT_GROUPS for c in cols if c["req"]]:
            if p.get(c["key"]) is None:
                rep.err("master", sc, "", c["header"], "Required field is empty")
        for c in [c for _, cols in S.PRODUCT_GROUPS for c in cols if c["list"]]:
            v = p.get(c["key"])
            if v is not None and v not in S.LISTS[c["list"]]:
                rep.err("master", sc, "", c["header"], f"'{v}' is not one of the dropdown options")
                p[c["key"]] = None
        for z in p["sizes"]:
            for c in [c for _, cols in S.SIZE_GROUPS for c in cols if c["req"]]:
                if z.get(c["key"]) is None:
                    rep.err("master", sc, z["sku"], c["header"], f"Required field is empty (Sizes row {z['_row']})")
        sanity_checks(sc, p, per_ch, rep)
    return prods, per_ch, general, size_tbl


def sanity_checks(sc, p, per_ch, rep):
    """Values that are allowed but look like a typing mistake."""
    meesho_brand = str((per_ch.get("Brand") or {}).get("meesho") or "")
    title = str(p.get("title") or "")
    if meesho_brand and "wish" not in meesho_brand.lower() and re.search(r"wish\s*karo", title + str(p.get("description") or ""), re.I):
        rep.info("meesho", sc, "", "Title/Description",
                 f"Meesho sells under '{meesho_brand}', so 'Wishkaro' in the Meesho title/description was replaced with '{meesho_brand}'.")
    if p.get("secondary_colour") and p.get("secondary_colour") == p.get("primary_colour"):
        p["secondary_colour"] = None
        rep.info("master", sc, "", "Secondary Colour", "Auto-fixed: same as Primary Colour, so it was left blank")
    main_imgs = [p.get(k) for k in ("img_front", "img_back", "img_side", "img_detail", "img_look")]
    if any("meeshosupplyassets" in str(v or "") for v in main_imgs):
        rep.warn("master", sc, "", "Image 1-5",
                 "Main image links are Meesho links. Flipkart/Myntra/Amazon may not accept them - use your own public links (Dropbox etc.)")
    zero = [z["size"] for z in p["sizes"] if z.get("stock") not in (None, "") and num(z.get("stock")) == 0]
    if zero:
        rep.warn("master", sc, "", "Stock", f"Stock is 0 for {', '.join(zero)} - these will list as out of stock")
    for z in p["sizes"]:
        wt = num(z.get("weight_g"))
        if isinstance(wt, (int, float)) and 0 < wt < 30:
            rep.warn("master", sc, z["sku"], "Weight (g)",
                     f"Weight is {wt} g - too light for a dress. Type it in grams (0.3 kg = 300)")
        sh, fl = z.get("shoulder"), z.get("front_length")
        if sh not in (None, "") and fl not in (None, "") and num(sh) == num(fl):
            z["shoulder"] = None
            rep.info("master", sc, z["sku"], "Across Shoulder (in)",
                     f"Auto-fixed: Shoulder ({sh}) was a copy of Front Length, so it was left blank")
        kw, auto = z.get("kids_weight"), V.MEESHO_KIDS_WEIGHT.get(z["size"])
        if kw and auto and kw != auto:
            lo = lambda s: float(str(s).split("-")[0])
            if abs(lo(kw) - lo(auto)) >= 5:
                z["kids_weight"] = None
                rep.info("meesho", sc, z["sku"], "Kids Weight",
                         f"Auto-fixed: {kw} kg is not possible for {z['size']}, used {auto} kg")


# ================================================================ helpers ===
def m(vmap, key, field, master_value, ch, rep, sc, sku, label):
    """Map a master value through a vocab dict."""
    if master_value is None:
        return None
    entry = vmap.get(master_value)
    if entry is None:
        return None
    v = entry.get(field)
    if v is None:
        rep.warn(ch, sc, sku, label, f"{CH_NAME[ch]} has no option for '{master_value}' - left blank, pick one manually")
    return v


def price_pair(p, z, ch):
    price = z.get(f"{ch}_price") if z.get(f"{ch}_price") is not None else p.get(f"{ch}_price")
    mrp = z.get(f"{ch}_mrp") if z.get(f"{ch}_mrp") is not None else p.get(f"{ch}_mrp")
    return price, mrp


def num(v):
    if v is None:
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() else round(f, 2)
    except (TypeError, ValueError):
        return v


def nearest(value, options):
    try:
        x = float(value)
    except (TypeError, ValueError):
        return None
    nums = [(abs(float(o) - x), o) for o in options if o.replace(".", "", 1).isdigit()]
    return min(nums)[1] if nums else None


def meesho_hsn(hsn):
    opts = ALLOWED["meesho"]["HSN ID"]
    for n in range(len(hsn), 3, -1):
        if hsn[:n] in opts:
            return hsn[:n]
    return None


def rebrand(text, brand):
    """Meesho sells under its own brand: swap our other brand names inside Meesho text."""
    if not text or not brand:
        return text
    return re.sub(r"wish\s*karo", brand, str(text), flags=re.I)


def short40(text):
    if not text or len(text) <= 40:
        return text
    cut = text[:40] if text[40] == " " else text[:40].rsplit(" ", 1)[0]
    words = cut.rstrip(" -,&").split(" ")
    while len(words) > 1 and words[-1].lower() in ("for", "with", "and", "of", "&", "in", "the", "a"):
        words.pop()
    return " ".join(words)


def stock_for(z, per_ch, ch):
    s = z.get("stock")
    return num(s) if s is not None else num(per_ch.get("Default Stock", {}).get(ch))


# ================================================================ builders ===
def build_meesho(prods, per_ch, gen, size_tbl, rep):
    ch = "meesho"
    rows, group_of = [], {}
    gap = num(gen.get("Meesho Returns Price Gap (Rs)"))
    for sc, p in prods.items():
        if not p["sizes"]:
            continue
        if p.get("meesho_price") is None or p.get("meesho_mrp") is None:
            if p["sizes"]:
                rep.info(ch, sc, "", "Price", "No Meesho price/MRP - not listed on Meesho")
            continue
        if p["parent_code"] not in group_of:
            group_of[p["parent_code"]] = f"Group {len(group_of) + 1:02d}"
        imgs = [p.get(f"meesho_img{i}") or p.get(k) for i, k in
                enumerate(["img_front", "img_back", "img_side", "img_detail"], start=1)]
        if not p.get("meesho_img1"):
            rep.warn(ch, sc, "", "Images", "No Meesho uploader links - used your normal image links. Meesho may reject non-Meesho links.")
        for z in p["sizes"]:
            sku = z["sku"]
            price, mrp = price_pair(p, z, ch)
            bust = z.get("bust") or z.get("chest")
            r = {
                "Product Name": rebrand(p["title"], per_ch["Brand"][ch]), "Variation": V.size_for(ch, z["size"]),
                "Meesho Price": price, "MRP": mrp,
                "Wrong/Defective Returns Price": (price - gap) if (gap and isinstance(price, (int, float)) and price > gap) else None,
                "GST %": p["gst"], "HSN ID": meesho_hsn(p["hsn"] or ""),
                "Net Weight (gms)": num(z.get("weight_g")), "Inventory": stock_for(z, per_ch, ch),
                "Country of Origin": per_ch["Country of Origin"][ch],
                "Manufacturer Details": per_ch["Manufacturer Details"][ch],
                "Packer Details": per_ch["Packer Details"][ch],
                "Importer Details": per_ch["Importer Details"][ch],
                "Fabric": m(V.FABRICS, None, ch, p["fabric"], ch, rep, sc, sku, "Fabric"),
                "Fit/ Shape": m(V.SHAPES, None, ch, p["shape"], ch, rep, sc, sku, "Fit/ Shape"),
                "Length": m(V.LENGTHS, None, ch, p["length"], ch, rep, sc, sku, "Length"),
                "Neck": m(V.NECKS, None, ch, p["neck"], ch, rep, sc, sku, "Neck"),
                "Net Quantity (N)": "Single",
                "Occasion": m(V.OCCASIONS, None, ch, p["occasion"], ch, rep, sc, sku, "Occasion"),
                "Sleeve Length": m(V.SLEEVE_LENGTHS, None, ch, p["sleeve_length"], ch, rep, sc, sku, "Sleeve Length"),
                "Surface Styling": m(V.SURFACE_STYLING, None, ch, p["surface"], ch, rep, sc, sku, "Surface Styling"),
                "Bust Size": nearest(bust, ALLOWED[ch]["Bust Size"]),
                "Kids Weight (In Kgs)": z.get("kids_weight") or size_tbl.get(z["size"]) or V.MEESHO_KIDS_WEIGHT[z["size"]],
                "Length Size": nearest(z.get("front_length"), ALLOWED[ch]["Length Size"]),
                "Image 1 (Front)": imgs[0], "Image 2": imgs[1], "Image 3": imgs[2], "Image 4": imgs[3],
                "Product ID / Style ID": sc, "SKU ID": sku,
                "Brand Name": per_ch["Brand"][ch], "Brand": per_ch["Brand"][ch],
                "Group ID": group_of[p["parent_code"]],
                "Product Description": rebrand(p["description"], per_ch["Brand"][ch]),
                "Add On": m(V.ADD_ONS, None, ch, p["add_on"], ch, rep, sc, sku, "Add On"),
                "Color": m(V.COLOURS, None, ch, p["primary_colour"], ch, rep, sc, sku, "Color"),
                "Pattern": m(V.PATTERNS, None, ch, p["pattern"], ch, rep, sc, sku, "Pattern"),
            }
            # Meesho's own sheet stores everything as text
            rows.append(({k: (str(v) if isinstance(v, (int, float)) else v) for k, v in r.items()}, sc, p["parent_code"]))
    return rows


def build_flipkart(prods, per_ch, gen, size_tbl, rep):
    ch = "flipkart"
    rows = []
    charge = num(gen.get("Flipkart Delivery Charge (Rs)")) or 0
    for sc, p in prods.items():
        if not p["sizes"]:
            continue
        if p.get("flipkart_price") is None or p.get("flipkart_mrp") is None:
            if p["sizes"]:
                rep.info(ch, sc, "", "Price", "No Flipkart price/MRP - not listed on Flipkart")
            continue
        bullets = [p.get(f"bullet{i}") for i in range(1, 6) if p.get(f"bullet{i}")]
        sstyle = p.get("sleeve_style") or ("No Sleeves" if p["sleeve_length"] == "Sleeveless" else "Regular Sleeves")
        hsn = int(p["hsn"]) if p.get("hsn") and str(p["hsn"]).isdigit() else p.get("hsn")
        for z in p["sizes"]:
            sku = z["sku"]
            price, mrp = price_pair(p, z, ch)
            fk_size = V.size_for(ch, z["size"])
            r = {
                "Seller SKU ID": sku, "Listing Status": gen.get("Flipkart Listing Status"),
                "MRP (INR)": num(mrp), "Your selling price (INR)": num(price),
                "Fullfilment by": gen.get("Flipkart Fulfilment By"),
                "Procurement type": gen.get("Flipkart Procurement Type"),
                "Procurement SLA (DAY)": num(gen.get("Flipkart Procurement SLA (days)")),
                "Stock": stock_for(z, per_ch, ch),
                "Shipping provider": gen.get("Flipkart Shipping Provider"),
                "Local delivery charge (INR)": charge, "Zonal delivery charge (INR)": charge,
                "National delivery charge (INR)": charge,
                "Length (CM)": num(z.get("pkg_l")), "Breadth (CM)": num(z.get("pkg_w")),
                "Weight (KG)": round(float(z["weight_g"]) / 1000, 3) if z.get("weight_g") else None,
                "Height (CM)": num(z.get("pkg_h")),
                "HSN": hsn, "Country Of Origin": per_ch["Country of Origin"][ch],
                "Manufacturer Details": per_ch["Manufacturer Details"][ch],
                "Packer Details": per_ch["Packer Details"][ch],
                "Importer Details": per_ch["Importer Details"][ch],
                "Tax Code": gen.get("Flipkart Tax Code"),
                "Brand": per_ch["Brand"][ch], "Style Code": sc, "Brand Color": p["brand_colour"],
                "Size": fk_size,
                "Type": m(V.SHAPES, None, ch, p["shape"], ch, rep, sc, sku, "Type"),
                "Occasion": m(V.OCCASIONS, None, ch, p["occasion"], ch, rep, sc, sku, "Occasion"),
                "Ideal For": "Baby Girls" if V.is_baby(z["size"]) else "Girls",
                "Primary Color": m(V.COLOURS, None, ch, p["primary_colour"], ch, rep, sc, sku, "Primary Color"),
                "Fabric": m(V.FABRICS, None, ch, p["fabric"], ch, rep, sc, sku, "Fabric"),
                "Sleeve Type": m(V.SLEEVE_LENGTHS, None, ch, p["sleeve_length"], ch, rep, sc, sku, "Sleeve Type"),
                "Suitable For": m(V.OCCASIONS, None, "flipkart_suitable", p["occasion"], ch, rep, sc, sku, "Suitable For"),
                "Neck Type": m(V.NECKS, None, ch, p["neck"], ch, rep, sc, sku, "Neck Type"),
                "Fabric Care": m(V.WASH_CARE, None, ch, p["care"], ch, rep, sc, sku, "Fabric Care"),
                "Character": p.get("character") or "No Character",
                "Length": m(V.LENGTHS, None, ch, p["length"], ch, rep, sc, sku, "Length"),
                "Number of Dresses": 1,
                "Fabric_Details": p.get("fabric_details") or p["fabric"],
                "Label Size": fk_size,
                "Main Image URL": p.get("img_front"), "Other Image URL 1": p.get("img_back"),
                "Other Image URL 2": p.get("img_side"), "Other Image URL 3": p.get("img_detail"),
                "Other Image URL 4": p.get("img_look"),
                "Group ID": p["parent_code"],
                "Description": p["description"], "Search Keywords": p.get("keywords"),
                "Key Features": "::".join(bullets) if bullets else None,
                "Secondary Color": m(V.COLOURS, None, ch, p.get("secondary_colour"), ch, rep, sc, sku, "Secondary Color"),
                "Closure": m(V.CLOSURES, None, ch, p["closure"], ch, rep, sc, sku, "Closure"),
                "Pattern": m(V.PATTERNS, None, ch, p["pattern"], ch, rep, sc, sku, "Pattern"),
                "Belt Included": V.ADD_ONS.get(p.get("add_on"), {}).get("flipkart_belt", "No"),
                "Surface Styling": m(V.SURFACE_STYLING, None, ch, p["surface"], ch, rep, sc, sku, "Surface Styling"),
                "Sleeve Style": m(V.SLEEVE_STYLES, None, ch, sstyle, ch, rep, sc, sku, "Sleeve Style"),
                "Pattern/Print Type": m(V.PATTERNS, None, "flipkart_print", p["pattern"], ch, rep, sc, sku, "Pattern/Print Type"),
            }
            rows.append((r, sc, p["parent_code"]))
    return rows


def build_myntra(prods, per_ch, gen, size_tbl, rep):
    ch = "myntra"
    rows = []
    group_ids = {}
    for sc, p in prods.items():
        if not p["sizes"]:
            continue
        if p.get("myntra_price") is None or p.get("myntra_mrp") is None:
            if p["sizes"]:
                rep.info(ch, sc, "", "Price", "No Myntra price/MRP - not listed on Myntra")
            continue
        art_name = p.get("short_name") or short40(p["title"])
        if art_name and len(art_name) > 40:
            rep.warn(ch, sc, "", "vendorArticleName", "Myntra Short Name is over 40 characters")
        elif not p.get("short_name") and art_name != p["title"]:
            rep.info(ch, sc, "", "vendorArticleName", f"Title cut to 40 chars for Myntra: '{art_name}'")
        sstyle = p.get("sleeve_style") or ("No Sleeves" if p["sleeve_length"] == "Sleeveless" else "Regular Sleeves")
        care = p["care"]
        for z in p["sizes"]:
            sku = z["sku"]
            price, mrp = price_pair(p, z, ch)
            gkey = (sc, mrp)  # Myntra: one styleGroupId per colour + MRP
            if gkey not in group_ids:
                group_ids[gkey] = len(group_ids) + 1
            my_size = V.size_for(ch, z["size"])
            r = {
                "styleGroupId": str(group_ids[gkey]),
                "vendorSkuCode": sku, "vendorArticleNumber": sku, "vendorArticleName": art_name,
                "brand": per_ch["Brand"][ch],
                "Manufacturer Name and Address with Pincode": per_ch["Manufacturer Details"][ch],
                "Packer Name and Address with Pincode": per_ch["Packer Details"][ch],
                "Importer Name and Address with Pincode": per_ch["Importer Details"][ch],
                "Country Of Origin": per_ch["Country of Origin"][ch],
                "articleType": "Dresses", "Brand Size": my_size, "Standard Size": my_size,
                "is Standard Size present on Label": "No",
                "Brand Colour (Remarks)": p["brand_colour"],
                "HSN": p["hsn"], "MRP": mrp, "ISP": price, "AgeGroup": "Kids-Girls",
                "Prominent Colour": m(V.COLOURS, None, ch, p["primary_colour"], ch, rep, sc, sku, "Prominent Colour"),
                "Second Prominent Colour": m(V.COLOURS, None, ch, p.get("secondary_colour"), ch, rep, sc, sku, "Second Prominent Colour"),
                "FashionType": gen.get("Myntra Fashion Type"),
                "Usage": m(V.OCCASIONS, None, "myntra_usage", p["occasion"], ch, rep, sc, sku, "Usage"),
                "Year": gen.get("Myntra Year"), "season": gen.get("Myntra Season"),
                "List View Name": p.get("list_view") or p["title"],
                "Product Details": p["description"],
                "materialCareDescription": f"{p.get('fabric_details') or p['fabric']}, {care}",
                "productDisplayName": p["title"],
                "Fabric": m(V.FABRICS, None, ch, p["fabric"], ch, rep, sc, sku, "Fabric"),
                "Occasion": m(V.OCCASIONS, None, ch, p["occasion"], ch, rep, sc, sku, "Occasion"),
                "Shape": m(V.SHAPES, None, ch, p["shape"], ch, rep, sc, sku, "Shape"),
                "Neck": m(V.NECKS, None, ch, p["neck"], ch, rep, sc, sku, "Neck"),
                "Pattern": m(V.PATTERNS, None, ch, p["pattern"], ch, rep, sc, sku, "Pattern"),
                "Length": m(V.LENGTHS, None, ch, p["length"], ch, rep, sc, sku, "Length"),
                "Sleeve Length": m(V.SLEEVE_LENGTHS, None, ch, p["sleeve_length"], ch, rep, sc, sku, "Sleeve Length"),
                "Knit or Woven": p["knit"],
                "Print or Pattern Type": m(V.PATTERNS, None, "myntra_print", p["pattern"], ch, rep, sc, sku, "Print or Pattern Type"),
                "Surface Styling": m(V.SURFACE_STYLING, None, ch, p["surface"], ch, rep, sc, sku, "Surface Styling"),
                "Main Trend": m(V.PATTERNS, None, "myntra_trend", p["pattern"], ch, rep, sc, sku, "Main Trend"),
                "Sleeve Styling": m(V.SLEEVE_STYLES, None, ch, sstyle, ch, rep, sc, sku, "Sleeve Styling"),
                "Transparency": p["transparency"],
                "Fabric Type": m(V.FABRICS, None, "myntra_type", p["fabric"], ch, rep, sc, sku, "Fabric Type"),
                "Lining": "Has a lining" if p.get("lining") == "Yes" else "NA",
                "Wash Care": m(V.WASH_CARE, None, ch, care, ch, rep, sc, sku, "Wash Care"),
                "Closure": m(V.CLOSURES, None, ch, p["closure"], ch, rep, sc, sku, "Closure"),
                "Add-Ons": m(V.ADD_ONS, None, ch, p["add_on"], ch, rep, sc, sku, "Add-Ons"),
                "Stitch": "Ready to Wear", "Sustainable": "Regular",
                "Character": ("NA" if str(p.get("character") or "").strip().lower() in ("no character", "none", "na", "no")
                              else p.get("character") if p.get("character") in ALLOWED[ch].get("Character", []) else None),
                "Number of Pockets": "NA", "Multipack Set": "NA", "Number of Items": "1",
                "Net Quantity Unit": "Piece", "Theme": p.get("theme"),
                "Package Contains": gen.get("Myntra Package Contains"), "Net Quantity": "1",
                "Bust ( Inches )": z.get("bust"), "Chest ( Inches )": z.get("chest"),
                "Front Length ( Inches )": z.get("front_length"), "Hips ( Inches )": z.get("hips"),
                "Waist ( Inches )": z.get("waist"), "Across Shoulder ( Inches )": z.get("shoulder"),
                "Front Image": p.get("img_front"), "Side Image": p.get("img_side"),
                "Back Image": p.get("img_back"), "Detail Angle": p.get("img_detail"),
                "Look Shot Image": p.get("img_look"), "Additional Image 1": p.get("img6"),
                "Additional Image 2": p.get("img7"),
            }
            if p.get("character") and r["Character"] is None:
                rep.warn(ch, sc, sku, "Character", f"'{p['character']}' is not in Myntra's character list - left blank")
            # Myntra's own sheet stores numbers as text
            rows.append(({k: (str(num(v)) if isinstance(v, (int, float)) else v) for k, v in r.items()}, sc, p["parent_code"]))
        for f_, label in (("img_side", "Side Image"), ("img_back", "Back Image"), ("img_detail", "Detail Angle"), ("img_look", "Look Shot Image")):
            if not p.get(f_):
                rep.warn(ch, sc, "", label, "Myntra marks this image as mandatory")
    return rows


# Amazon column keys (row 5 of the template, without the [marketplace_id=...] parts)
AMZ_SKU = "contribution_sku#1.value"
AMZ_PARENTAGE = "parentage_level#1.value"
AMZ_PRICE = "purchasable_offer#1.our_price#1.schedule#1.value_with_tax"
AMZ_MRP = "purchasable_offer#1.maximum_retail_price#1.schedule#1.value_with_tax"


def amz_pick(key, value):
    """Exact spelling of value in Amazon's dropdown for key (ignores case and spaces). None if not there."""
    if value in (None, ""):
        return None
    opts = ALLOWED["amazon"].get(key) or []
    norm = lambda s: re.sub(r"[\s\-]+", "", str(s)).lower()
    return next((o for o in opts if norm(o) == norm(value)), None)


def amz_multi(base, key_suffix, values, limit):
    """Fill repeated Amazon columns (bullet_point#1..#5 etc.) from a list."""
    return {f"{base}#{i}{key_suffix}": v for i, v in enumerate(values[:limit], start=1)}


def amz_setting_list(gen, label, key, limit, rep):
    vals = []
    for part in str(gen.get(label, S.GENERAL_DEFAULTS.get(label)) or "").split(","):
        v = amz_pick(key, part.strip())
        if v and v not in vals:
            vals.append(v)
        elif part.strip() and not v:
            rep.warn("amazon", "", "", label, f"'{part.strip()}' is not an Amazon option - ignored")
    return vals[:limit]


def parent_title(titles):
    """Words shared by every colour's title, e.g. '... Cotton Red & Black Checked ...' and
    '... Cotton Green & Blue Checked ...' -> '... Cotton Checked ...'."""
    words = [t.split() for t in titles]
    if len(words) < 2:
        return titles[0] if titles else None
    shared = set.intersection(*(set(w) for w in words[1:]))
    first = words[0]
    keep = [w in shared for w in first]
    for i, w in enumerate(first):                    # drop '&' / 'and' that joined dropped colour words
        if keep[i] and w.lower() in ("&", "and", "/") and ((i and not keep[i - 1]) or (i + 1 < len(first) and not keep[i + 1])):
            keep[i] = False
    return " ".join(w for w, k in zip(first, keep) if k) or titles[0]


def amz_date(v, rep, label):
    if v in (None, ""):
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            pass
    rep.err("amazon", "", "", label, f"'{v}' is not a date - sale dates left blank. Type it as YYYY-MM-DD")
    return None


def build_amazon(prods, per_ch, gen, size_tbl, rep):
    ch = "amazon"
    rows = []
    G = lambda label: gen.get(label, S.GENERAL_DEFAULTS.get(label))
    start = amz_date(G("Amazon Sale Start Date"), rep, "Amazon Sale Start Date")
    end = amz_date(G("Amazon Sale End Date"), rep, "Amazon Sale End Date")
    if start and not end:
        end = start.replace(year=start.year + 2, day=min(start.day, 28) if start.month == 2 else start.day)
    pad = num(G("Amazon Package Padding (cm)")) or 0
    min_gap, max_gap = num(G("Amazon Min Price Gap (Rs)")), num(G("Amazon Max Price Gap (Rs)"))
    brand_in = per_ch["Brand"][ch]
    brand = amz_pick("brand#1.value", brand_in) or brand_in
    if brand == brand_in and ALLOWED[ch].get("brand#1.value") and brand not in ALLOWED[ch]["brand#1.value"]:
        rep.err(ch, "", "", "Brand", f"Amazon brand '{brand_in}' is not approved in your Amazon template "
                                     f"(allowed: {', '.join(ALLOWED[ch]['brand#1.value'])})")
    mfr = per_ch["Manufacturer Details"][ch]
    packer = per_ch["Packer Details"][ch] or mfr
    importer = per_ch["Importer Details"][ch] or mfr
    country = per_ch["Country of Origin"][ch]

    # browse nodes: setting holds numbers, the template wants its own "Path (number)" labels
    node_labels = ALLOWED[ch].get("recommended_browse_nodes#1.value", [])
    nodes = []
    for node in re.findall(r"\d{6,}", str(G("Amazon Browse Node") or "")):
        label = next((l for l in node_labels if l.endswith(f"({node})")), None)
        if not label:
            rep.warn(ch, "", "", "Amazon Browse Node", f"Node {node} is not in the Amazon template's list - used the number as is")
        nodes.append(label or node)
    features_base = amz_setting_list(gen, "Amazon Special Features", "special_feature#1.value", 5, rep)
    occasions = amz_setting_list(gen, "Amazon Occasions", "occasion_type#1.value", 5, rep)
    holidays = amz_setting_list(gen, "Amazon Holidays", "holiday_type#1.value", 3, rep)
    life_events = amz_setting_list(gen, "Amazon Life Events", "life_event#1.value", 3, rep)
    seasons = amz_setting_list(gen, "Amazon Seasons", "seasons#1.value", 5, rep)
    id_type = amz_pick("amzn1.volt.ca.product_id_type", G("Amazon Product ID Type")) or "GTIN Exempt"
    fulfil = amz_pick("fulfillment_availability#1.fulfillment_channel_code", G("Amazon Fulfilment")) \
        or "Fulfillment by Merchant (Default)"
    handling = num(G("Amazon Handling Time (days)"))

    # one parent title and one price ceiling per design (all its colours and sizes)
    design_titles, design_top = defaultdict(list), defaultdict(float)
    for sc, p in prods.items():
        if p["sizes"] and p.get("amazon_price") is not None and p.get("amazon_mrp") is not None:
            design_titles[p["parent_code"]].append(p["title"])
            for z in p["sizes"]:
                pr = num(price_pair(p, z, ch)[0])
                if isinstance(pr, (int, float)):
                    design_top[p["parent_code"]] = max(design_top[p["parent_code"]], pr)

    parents_done = set()
    for sc, p in prods.items():
        if not p["sizes"]:
            continue
        if p.get("amazon_price") is None or p.get("amazon_mrp") is None:
            rep.info(ch, sc, "", "Price", "No Amazon price/MRP - not listed on Amazon")
            continue
        features = list(features_base)
        if p.get("transparency") == "Opaque":
            features.append("Not See-Through")
        if p.get("lining") == "Yes":
            features.append("Lined")
        features = list(dict.fromkeys(features))[:5]
        sleeve_type = V.SLEEVE_STYLES.get(p.get("sleeve_style") or "", {}).get("amazon")
        stretch = ("Medium Stretch" if p["fabric"] == "Cotton Lycra"
                   else "Low Stretch" if p.get("knit") in ("Knitted", "Knitted and Woven") else "No Stretch")
        lining = None
        if p.get("lining") == "Yes":
            lining = amz_pick("lining_description#1.value", V.FABRICS.get(p["fabric"], {}).get("amazon")) or "Polyester"
        # same on the parent row and every size, as in the listing team's own Amazon file
        common = {
            "product_type#1.value": "DRESS", "::record_action": "Create or Replace (Full Update)",
            "variation_theme#1.name": "SIZE/COLOR",
            "item_name#1.value": p["title"], "brand#1.value": brand, "manufacturer#1.value": mfr,
            "product_description#1.value": p["description"],
            "generic_keyword#1.value": p.get("keywords"),
            "lifestyle#1.value": m(V.OCCASIONS, None, ch, p["occasion"], ch, rep, sc, "", "Lifestyle"),
            "target_gender#1.value": "Female",
            "fabric_type#1.value": p.get("fabric_details") or p["fabric"],
            "lining_description#1.value": lining,
            "item_type_name#1.value": V.OCCASIONS.get(p["occasion"] or "", {}).get("amazon_type"),
            "part_number#1.value": p["parent_code"],
            "theme#1.value": amz_pick("theme#1.value", p.get("theme")),
            "care_instructions#1.value": m(V.WASH_CARE, None, ch, p["care"], ch, rep, sc, "", "Care Instructions"),
            "rtip_manufacturer_contact_information#1.value": mfr,
            "back#1.style#1.value": V.CLOSURES.get(p["closure"] or "", {}).get("amazon_back"),
            "external_product_information#1.entity": "HSN Code",
            "external_product_information#1.value": p["hsn"],
            "importer_contact_information#1.value": importer,
            "packer_contact_information#1.value": packer,
            "apparel_silhouette#1.value": V.SHAPES.get(p["shape"] or "", {}).get("amazon_silhouette"),
            "sleeve#1.type#1.value": sleeve_type,
            "closure#1.type#1.value": m(V.CLOSURES, None, ch, p["closure"], ch, rep, sc, "", "Closure Type"),
            "apparel_fabric_weight_class#1.value": "Lightweight",
            "garment_size_country#1.value": "India",
            "apparel_fabric_stretch#1.value": stretch,
            "fit_to_size_sentiment#1.value": "Fits True To Size",
            "formality_level#1.value": V.OCCASIONS.get(p["occasion"] or "", {}).get("amazon_formality"),
            "country_of_origin#1.value": country,
            "batteries_required#1.value": "No",
        }
        common.update(amz_multi("recommended_browse_nodes", ".value", nodes, 5))
        common.update(amz_multi("bullet_point", ".value", [p.get(f"bullet{i}") for i in range(1, 6) if p.get(f"bullet{i}")], 5))
        common.update(amz_multi("special_feature", ".value", features, 5))
        common.update(amz_multi("occasion_type", ".value", occasions, 5))
        common.update(amz_multi("seasons", ".value", seasons, 5))
        common.update(amz_multi("holiday_type", ".value", holidays, 3))
        common.update(amz_multi("life_event", ".value", life_events, 3))

        if p["parent_code"] not in parents_done:
            parents_done.add(p["parent_code"])
            parent = dict(common)
            parent.update({AMZ_SKU: p["parent_code"], AMZ_PARENTAGE: "Parent",
                           "item_name#1.value": parent_title(design_titles[p["parent_code"]])})
            rows.append((parent, sc, p["parent_code"]))

        others = [p.get(k) for k in ("img_back", "img_side", "img_detail", "img_look", "img6", "img7", "img8")]
        images = {"main_product_image_locator#1.media_location": p.get("img_front")}
        images.update({f"other_product_image_locator_{i}#1.media_location": u
                       for i, u in enumerate([o for o in others if o], start=1)})
        neck = V.NECKS.get(p["neck"], {})
        for z in p["sizes"]:
            sku = z["sku"]
            price, mrp = num(price_pair(p, z, ch)[0]), num(price_pair(p, z, ch)[1])
            lo, hi = V.amazon_size_range(z["size"])
            L_, W_, H_ = num(z.get("pkg_l")), num(z.get("pkg_w")), num(z.get("pkg_h"))
            wt = num(z.get("weight_g"))
            is_num = lambda v: isinstance(v, (int, float))
            r = dict(common)
            r.update(images)
            r.update({
                AMZ_SKU: sku, AMZ_PARENTAGE: "Child",
                "child_parent_sku_relationship#1.parent_sku": p["parent_code"],
                "amzn1.volt.ca.product_id_type": id_type,
                "model_number#1.value": p["parent_code"],
                "model_name#1.value": G("Amazon Model Name") or p["title"],
                "collar_style#1.value": neck.get("amazon_collar"),
                "style#1.value": m(V.SHAPES, None, ch, p["shape"], ch, rep, sc, sku, "Style"),
                "department#1.value": "baby-girls" if V.is_baby(z["size"]) else "girls",
                "age_range_description#1.value": "Baby" if V.is_baby(z["size"]) else "Kid",
                "apparel_size#1.size_system": "IN", "apparel_size#1.size_class": "Age",
                "apparel_size#1.size": lo, "apparel_size#1.size_to": hi, "apparel_size#1.body_type": "Regular",
                "material#1.value": m(V.FABRICS, None, ch, p["fabric"], ch, rep, sc, sku, "Material"),
                "number_of_items#1.value": 1,
                "color#1.standardized_values#1": m(V.COLOURS, None, ch, p["primary_colour"], ch, rep, sc, sku, "Color Map"),
                "color#1.value": p["brand_colour"],
                "item_length_description#1.value": m(V.LENGTHS, None, ch, p["length"], ch, rep, sc, sku, "Item Length Description"),
                "fit_type#1.value": "Regular",
                "pattern#1.value": m(V.PATTERNS, None, ch, p["pattern"], ch, rep, sc, sku, "Pattern"),
                "unit_count#1.value": 1, "unit_count#1.type.value": "count",
                "neck#1.neck_style#1.value": m(V.NECKS, None, ch, p["neck"], ch, rep, sc, sku, "Neck Style"),
                "sleeve#1.length_description#1.value": m(V.SLEEVE_LENGTHS, None, ch, p["sleeve_length"], ch, rep, sc, sku, "Sleeve Length"),
                "item_weight#1.value": wt, "item_weight#1.unit": "Grams" if wt is not None else None,
                "condition_type#1.value": "New",
                "fulfillment_availability#1.fulfillment_channel_code": fulfil,
                "fulfillment_availability#1.quantity": stock_for(z, per_ch, ch),
                "fulfillment_availability#1.lead_time_to_ship_max_days": handling,
                AMZ_PRICE: price, AMZ_MRP: mrp,
                "purchasable_offer#1.minimum_seller_allowed_price#1.schedule#1.value_with_tax":
                    (price - min_gap) if is_num(price) and is_num(min_gap) and price > min_gap else None,
                "purchasable_offer#1.maximum_seller_allowed_price#1.schedule#1.value_with_tax":
                    num(min(max(design_top[p["parent_code"]], price) + max_gap, mrp))
                    if is_num(price) and is_num(max_gap) and is_num(mrp) else None,
                "purchasable_offer#1.discounted_price#1.schedule#1.value_with_tax": price,
                "purchasable_offer#1.discounted_price#1.schedule#1.start_at": start.isoformat() if start else None,
                "purchasable_offer#1.discounted_price#1.schedule#1.end_at": end.isoformat() if end else None,
                "item_dimensions#1.length.value": L_, "item_dimensions#1.width.value": W_, "item_dimensions#1.height.value": H_,
                "item_package_dimensions#1.length.value": (L_ + pad) if is_num(L_) else None,
                "item_package_dimensions#1.width.value": (W_ + pad) if is_num(W_) else None,
                "item_package_dimensions#1.height.value": H_,
                "item_package_weight#1.value": wt, "item_package_weight#1.unit": "Grams" if wt is not None else None,
            })
            for dim in ("item_dimensions", "item_package_dimensions"):
                for side in ("length", "width", "height"):
                    if r.get(f"{dim}#1.{side}.value") is not None:
                        r[f"{dim}#1.{side}.unit"] = "Centimetres"
            rows.append((r, sc, p["parent_code"]))
    return rows


BUILDERS = {"meesho": build_meesho, "flipkart": build_flipkart, "myntra": build_myntra, "amazon": build_amazon}

# Which template columns are dropdowns -> allowed list name in allowed_values.json
DROPDOWN_CHECK = {
    "meesho": {"Variation": "Variation", "GST %": "GST %", "HSN ID": "HSN ID", "Fabric": "Fabric",
               "Fit/ Shape": "Fit/ Shape", "Length": "Length", "Neck": "Neck", "Occasion": "Occasion",
               "Sleeve Length": "Sleeve Length", "Surface Styling": "Surface Styling", "Bust Size": "Bust Size",
               "Kids Weight (In Kgs)": "Kids Weight (In Kgs)", "Length Size": "Length Size",
               "Add On": "Add On", "Brand": "Brand", "Color": "Color", "Pattern": "Pattern",
               "Country of Origin": "Country of Origin"},
    "flipkart": {"Size": "Size", "Type": "Type", "Occasion": "Occasion", "Ideal For": "Ideal For",
                 "Primary Color": "Primary Color", "Fabric": "Fabric", "Sleeve Type": "Sleeve Type",
                 "Suitable For": "Suitable For", "Neck Type": "Neck Type", "Fabric Care": "Fabric Care",
                 "Length": "Length", "Secondary Color": "Secondary Color", "Closure": "Closure",
                 "Pattern": "Pattern", "Surface Styling": "Surface Styling", "Sleeve Style": "Sleeve Style",
                 "Pattern/Print Type": "Pattern/Print Type", "Country Of Origin": "Country Of Origin"},
    "myntra": {k: k for k in ["brand", "Country Of Origin", "articleType", "Brand Size", "Standard Size",
               "is Standard Size present on Label", "AgeGroup", "Prominent Colour", "Second Prominent Colour",
               "FashionType", "Usage", "Year", "season", "Fabric", "Occasion", "Shape", "Neck", "Pattern",
               "Length", "Sleeve Length", "Knit or Woven", "Print or Pattern Type", "Surface Styling",
               "Main Trend", "Sleeve Styling", "Transparency", "Fabric Type", "Lining", "Wash Care",
               "Closure", "Add-Ons", "Stitch", "Sustainable", "Number of Pockets", "Multipack Set",
               "Number of Items", "Net Quantity Unit", "Theme"]},
    "amazon": None,  # every column that has a list in the template's "Dropdown Lists" sheet
}
REQUIRED = {
    "meesho": None,  # read from the template's "* Compulsory Field" row
    "flipkart": ["Seller SKU ID", "MRP (INR)", "Your selling price (INR)", "HSN", "Brand", "Style Code",
                 "Size", "Type", "Ideal For", "Primary Color", "Fabric", "Main Image URL", "Description"],
    "myntra": ["vendorSkuCode", "vendorArticleNumber", "vendorArticleName", "brand", "Brand Size", "HSN",
               "MRP", "ISP", "Prominent Colour", "Fabric", "Shape", "Neck", "Length", "Sleeve Length",
               "Chest ( Inches )", "Front Length ( Inches )", "Front Image"],
    "amazon": [AMZ_SKU, "brand#1.value", "item_name#1.value", AMZ_PRICE, AMZ_MRP,
               "main_product_image_locator#1.media_location", "apparel_size#1.size", "apparel_size#1.size_to",
               "color#1.standardized_values#1", "material#1.value", "department#1.value",
               "item_package_weight#1.value"],
}


# ================================================================ templates ===
def template_headers(ch):
    t = TEMPLATES[ch]
    path = os.path.join(TPL, t["file"])
    if t["ext"] == "xls":
        s = xlrd.open_workbook(path).sheet_by_name(t["sheet"])
        raw = s.row_values(t["header_row"] - 1)
        extra = None
    else:
        ws = openpyxl.load_workbook(path, read_only=True)[t["sheet"]]
        raw = list(ws.iter_rows(min_row=t["header_row"], max_row=t["header_row"], values_only=True))[0]
        extra = list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0] if ch == "meesho" else None
    heads = {}
    compulsory = []
    for i, h in enumerate(raw, start=1):
        if h in (None, ""):
            continue
        h = str(h)
        if ch == "amazon":
            h = V.amazon_key(h)
        if ch == "meesho":
            parts = [x.strip() for x in h.split("\n\n") if x.strip()]
            h = parts[0] if parts else h
            if extra and i - 1 < len(extra) and extra[i - 1] and "Compulsory" in str(extra[i - 1]):
                compulsory.append(h)
        heads.setdefault(h.strip(), i)
    return heads, compulsory


def check_rows(ch, rows, heads, compulsory, rep):
    allowed = ALLOWED[ch]
    required = compulsory if REQUIRED[ch] is None else REQUIRED[ch]
    for r, sc, _ in rows:
        sku = r.get("SKU ID") or r.get("Seller SKU ID") or r.get("vendorSkuCode") or r.get(AMZ_SKU)
        is_amazon_parent = ch == "amazon" and r.get(AMZ_PARENTAGE) == "Parent"
        for k in r:
            if k not in heads and r[k] not in (None, ""):
                rep.err(ch, sc, sku, k, "Column not found in template (template changed?) - value dropped")
        for k in ([] if is_amazon_parent else required):
            if k in ("Wrong/Defective Returns Price",):
                continue
            if r.get(k) in (None, ""):
                rep.err(ch, sc, sku, k, "Mandatory on this platform but empty")
        checks = DROPDOWN_CHECK[ch] if DROPDOWN_CHECK[ch] is not None else \
            {k: k for k in r if k in allowed and not isinstance(r[k], (int, float))}
        for col, lst in checks.items():
            v = r.get(col)
            opts = allowed.get(lst)
            if v not in (None, "") and opts and str(v) not in opts:
                rep.err(ch, sc, sku, col, f"'{v}' is not in {CH_NAME[ch]}'s dropdown")
        if is_amazon_parent:
            continue
        # price sanity
        pr = {"meesho": ("Meesho Price", "MRP"), "flipkart": ("Your selling price (INR)", "MRP (INR)"),
              "myntra": ("ISP", "MRP"), "amazon": (AMZ_PRICE, AMZ_MRP)}[ch]
        try:
            if float(r[pr[0]]) > float(r[pr[1]]):
                rep.err(ch, sc, sku, "Price", f"Selling price {r[pr[0]]} is higher than MRP {r[pr[1]]}")
        except (TypeError, ValueError, KeyError):
            pass
        if ch == "flipkart" and isinstance(r.get("Seller SKU ID"), str) and len(r["Seller SKU ID"]) > 64:
            rep.err(ch, sc, sku, "Seller SKU ID", "Longer than 64 characters")


def chunk(rows, max_rows, max_groups=None):
    """Split rows into files without breaking a design (parent) across files."""
    blocks = []
    for row in rows:
        if blocks and blocks[-1][0] == row[2]:
            blocks[-1][1].append(row)
        else:
            blocks.append((row[2], [row]))
    out, cur, groups = [], [], 0
    for _, block in blocks:
        too_many_rows = cur and len(cur) + len(block) > max_rows
        too_many_groups = max_groups and groups >= max_groups
        if too_many_rows or too_many_groups:
            out.append(cur); cur, groups = [], 0
        cur.extend(block); groups += 1
    if cur:
        out.append(cur)
    return out


def regroup_meesho(rows):
    """Group IDs restart at Group 01 in every Meesho file."""
    seen = {}
    for r, sc, parent in rows:
        if parent not in seen:
            seen[parent] = f"Group {len(seen) + 1:02d}"
        r["Group ID"] = seen[parent]


# ================================================================== report ===
def write_report(path, rep, summary):
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Platform", "Products", "Rows (SKUs)", "File(s)"])
    for row in summary:
        ws.append(row)
    ws.append([])
    counts = defaultdict(int)
    for lvl, *_ in rep.rows:
        counts[lvl] += 1
    ws.append(["ERROR (fix before upload)", counts["ERROR"]])
    ws.append(["CHECK (look at it)", counts["CHECK"]])
    ws.append(["INFO", counts["INFO"]])
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for col, w in zip("ABCD", (26, 12, 14, 70)):
        ws.column_dimensions[col].width = w
    ws2 = wb.create_sheet("Issues")
    ws2.append(["Level", "Platform", "Style Code", "SKU", "Field", "What to do"])
    order = {"ERROR": 0, "CHECK": 1, "INFO": 2}
    fills = {"ERROR": "F8CBAD", "CHECK": "FFE699", "INFO": "DDEBF7"}
    # one line per (problem, product) instead of one per size
    grouped = OrderedDict()
    for lvl, plat, style, sku, field, msg in rep.rows:
        grouped.setdefault((lvl, plat, style, field, msg), [])
        if sku and sku not in grouped[(lvl, plat, style, field, msg)]:
            grouped[(lvl, plat, style, field, msg)].append(sku)
    lines = []
    for (lvl, plat, style, field, msg), skus in grouped.items():
        if len(skus) > 2:
            sku_txt = f"{len(skus)} sizes"
        else:
            sku_txt = ", ".join(skus)
        lines.append((lvl, plat, style, sku_txt, field, msg))
    for row in sorted(lines, key=lambda x: (order[x[0]], x[1] != "Master sheet", x[1], x[2], x[4])):
        ws2.append(list(row))
        for c in ws2[ws2.max_row]:
            c.fill = PatternFill("solid", fgColor=fills[row[0]])
            c.font = Font(name="Arial", size=10)
    for c in ws2[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for col, w in zip("ABCDEF", (9, 11, 18, 28, 26, 90)):
        ws2.column_dimensions[col].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions
    wb.save(path)


# ==================================================================== main ===
def run(master, out_dir):
    os.makedirs(out_dir, exist_ok=True)
    rep = Report()
    prods, per_ch, gen, size_tbl = load_master(master, rep)
    stamp = dt.datetime.now().strftime("%Y%m%d_%H%M")
    summary = []
    for ch, builder in BUILDERS.items():
        rows = builder(prods, per_ch, gen, size_tbl, rep)
        if not rows:
            summary.append([CH_NAME[ch], 0, 0, "- nothing to list"])
            continue
        heads, compulsory = template_headers(ch)
        check_rows(ch, rows, heads, compulsory, rep)
        max_rows = num(per_ch.get("Max Rows per File", {}).get(ch))
        if not isinstance(max_rows, (int, float)) or max_rows < 1:
            if max_rows not in (None, ""):
                rep.err(ch, "", "", "Max Rows per File", f"'{max_rows}' is not a number - wrote one file")
            max_rows = 10**6
        parts = chunk(rows, int(max_rows), 30 if ch == "meesho" else None)
        files = []
        for i, part in enumerate(parts, start=1):
            if ch == "meesho":
                regroup_meesho(part)
            t = TEMPLATES[ch]
            suffix = f"_part{i}" if len(parts) > 1 else ""
            out = os.path.join(out_dir, f"{CH_NAME[ch]}_{stamp}{suffix}.{t['ext']}")
            cells = [{heads[k]: v for k, v in r.items() if k in heads and v not in (None, "")} for r, _, _ in part]
            src = os.path.join(TPL, t["file"])
            if t["ext"] == "xls":
                writers.fill_xls(src, out, t["sheet"], t["data_row"], cells)
            else:
                writers.fill_ooxml(src, out, t["sheet"], t["data_row"], cells)
            files.append(os.path.basename(out))
        styles = len({sc for _, sc, _ in rows})
        summary.append([CH_NAME[ch], styles, len(rows), ", ".join(files)])
    report_path = os.path.join(out_dir, f"Listing_Report_{stamp}.xlsx")
    write_report(report_path, rep, summary)
    return summary, rep, report_path


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    master = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.getcwd(), "output")
    summary, rep, report = run(master, out_dir)
    for s in summary:
        print(f"{s[0]:9} {s[1]:>4} products {s[2]:>5} rows  -> {s[3]}")
    lv = defaultdict(int)
    for r in rep.rows:
        lv[r[0]] += 1
    print(f"Report: {report}  (errors {lv['ERROR']}, checks {lv['CHECK']}, info {lv['INFO']})")
