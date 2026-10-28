"""
Run this after you replace a file in templates/ with a newer one downloaded
from Meesho / Flipkart / Myntra / Amazon. It re-reads every dropdown list into
config/allowed_values.json and re-checks vocab.py against it.

    python refresh_templates.py
"""
import json
import os
import re
import warnings
import zipfile

import openpyxl
import xlrd
from openpyxl.utils import column_index_from_string as ci

warnings.filterwarnings("ignore")
HERE = os.path.dirname(os.path.abspath(__file__))
T = os.path.join(HERE, "templates")


def _rng(rows, ref):
    m = re.match(r"\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?", ref)
    c, r1, r2 = ci(m.group(1)), int(m.group(2)), int(m.group(4) or m.group(2))
    return [str(rows[r - 1][c - 1]) for r in range(r1, r2 + 1)
            if r - 1 < len(rows) and c - 1 < len(rows[r - 1]) and rows[r - 1][c - 1] not in (None, "")]


def meesho():
    ws = openpyxl.load_workbook(os.path.join(T, "meesho.xlsx"), read_only=True)["Validation Sheet"]
    rows = list(ws.iter_rows(values_only=True))
    out = {}
    for j, h in enumerate(rows[0]):
        if h:
            vals = [str(r[j]) for r in rows[2:] if j < len(r) and r[j] not in (None, "")]
            if vals:
                out[str(h)] = vals
    return out


def flipkart():
    wb = xlrd.open_workbook(os.path.join(T, "flipkart.xls"))
    s = wb.sheet_by_name("Index")
    out = {}
    for c, h in enumerate(s.row_values(1)):
        if str(h).strip() and c > 0:
            out[str(h).strip()] = [str(s.cell_value(r, c)).strip() for r in range(2, s.nrows) if str(s.cell_value(r, c)).strip()]
    fill = wb.sheet_by_name("kids_dress").row_values(0)
    for name in wb.sheet_names():
        m = re.match(r"DropDownValuesForColumn(\d+)", name)
        if m:  # these sheets hold the list for template column N (0-based)
            col_name = fill[int(m.group(1))]
            out[col_name] = [str(x).strip() for x in wb.sheet_by_name(name).col_values(0) if str(x).strip()]
    return out


def myntra():
    path = os.path.join(T, "myntra.xlsx")
    x = zipfile.ZipFile(path).read("xl/worksheets/sheet2.xml").decode()
    wb = openpyxl.load_workbook(path, read_only=True)
    md = list(wb["masterdata"].iter_rows(values_only=True))
    hdr = list(wb["Dresses"].iter_rows(min_row=3, max_row=3, values_only=True))[0]
    out = {}
    for m in re.finditer(r"<x14:dataValidation\b.*?<xm:f>masterdata!(.*?)</xm:f>.*?<xm:sqref>(.*?)</xm:sqref>", x, re.S):
        col = re.match(r"([A-Z]+)", m.group(2)).group(1)
        vals = _rng(md, m.group(1))
        if vals:
            out[hdr[ci(col) - 1]] = vals
    return out


def amazon():
    """New-style Amazon template: 'Dropdown Lists' row 3 holds the attribute key of
    each list, the allowed values sit below it. Stored under vocab.amazon_key()."""
    import vocab
    path = os.path.join(T, "amazon.xlsm")
    dl = list(openpyxl.load_workbook(path, read_only=True)["Dropdown Lists"].iter_rows(values_only=True))
    out = {}
    for j, h in enumerate(dl[2]):
        if not h or not ("#" in str(h) or str(h).startswith("::") or str(h).startswith("amzn1.")):
            continue
        vals = out.setdefault(vocab.amazon_key(str(h)), [])
        for r in dl[3:]:
            if j < len(r) and r[j] not in (None, "") and str(r[j]) not in vals:
                vals.append(str(r[j]))
    return out


if __name__ == "__main__":
    allowed = {"meesho": meesho(), "flipkart": flipkart(), "myntra": myntra(), "amazon": amazon()}
    for ch, d in allowed.items():
        print(f"{ch:9} {len(d)} dropdown lists")
    with open(os.path.join(HERE, "config", "allowed_values.json"), "w", encoding="utf-8") as fh:
        json.dump(allowed, fh, ensure_ascii=False, indent=0)
    import vocab
    probs = vocab.self_check(allowed)
    print("\n".join(probs) if probs else "vocab.py still matches every template. Done.")
