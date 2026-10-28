"""
Builds the blank master sheet the listing person fills in.
    python build_master.py              -> Wishkaro_Master_Listing.xlsx
Re-run after editing vocab.py or master_spec.py so dropdowns stay in sync.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule

import master_spec as S
import vocab

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
GROUP_FILLS = ["2E75B6", "7030A0", "D81B60", "00897B", "B7791F", "2E7D32", "E64A19", "5C6BC0"]
# platform colours (headers of platform-specific columns use these)
PLATFORM_COLOURS = {"Meesho": "9F2089", "Flipkart": "2874F0", "Myntra": "FF3F6C", "Amazon": "F08804"}


def tint(hex_colour, amount=0.88):
    """Mix a colour with white. amount=0.88 -> very light version of the colour."""
    r, g, b = (int(hex_colour[i:i + 2], 16) for i in (0, 2, 4))
    mix = lambda v: int(v + (255 - v) * amount)
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


def platform_of(header):
    for name, colour in PLATFORM_COLOURS.items():
        if header.startswith(name):
            return colour
    return None
REQ_FILL = PatternFill("solid", fgColor="FFF2CC")     # yellow = must fill
OPT_FILL = PatternFill("solid", fgColor="F2F2F2")
EX_FILL = PatternFill("solid", fgColor="E2EFDA")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATA_ROWS = 600                                        # rows with dropdowns

EXAMPLE_PRODUCT = dict(
    style_code="EXAMPLE_fe3330nb", parent_code="EXAMPLE_fe3330",
    title="Girls Tri-Color Pleated Dress", short_name="Girls Tri-Color Pleated Dress",
    description="Dress your little patriot in this charming tri-colour dress with soft, breathable fabric and a knee-length flared fit.",
    bullet1="Material: Lightweight net fabric with soft cotton lining for all-day comfort",
    bullet2="Design: Tricolour pleats with a clean flare",
    bullet3="Styling Tip: Pair with white shoes and a hairband",
    bullet4="Fit: Regular fit with a knee-length hemline for easy movement",
    keywords="republic day dress, tricolor dress, kids dress, girls dress",
    brand_colour="Multicolor", primary_colour="Multicolor", fabric="Net",
    fabric_details="Net with cotton lining", shape="Fit and Flare", length="Knee Length",
    neck="Round Neck", sleeve_length="Sleeveless", pattern="Solid", occasion="Party",
    surface="None", closure="Zip", care="Hand Wash", lining="Yes", knit="Woven",
    transparency="Opaque", add_on="Hairband", theme="NA", hsn="62041200", gst="5",
    meesho_price=699, meesho_mrp=999, flipkart_price=799, flipkart_mrp=1999,
    myntra_price=1149, myntra_mrp=2899, amazon_price=799, amazon_mrp=1999,
    img_front="https://www.dropbox.com/scl/fi/.../8680_1.jpg?dl=0",
    img_back="https://www.dropbox.com/scl/fi/.../8680_3.jpg?dl=0",
    img_side="https://www.dropbox.com/scl/fi/.../8680_4.jpg?dl=0",
)
EXAMPLE_SIZES = [
    dict(style_code="EXAMPLE_fe3330nb", size="3-6 Months", chest=20, front_length=24.5,
         bust=21.7, waist=8.8, hips=21.5, pkg_l=26, pkg_w=22, pkg_h=2, weight_g=220),
    dict(style_code="EXAMPLE_fe3330nb", size="2-3 Years", chest=21, front_length=25,
         bust=22.2, waist=9.1, hips=22, pkg_l=28, pkg_w=23, pkg_h=2, weight_g=280,
         amazon_price=899, amazon_mrp=2149),
]


def f(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def write_lists(wb):
    ws = wb.create_sheet("Lists")
    refs = {}
    for i, (name, values) in enumerate(S.LISTS.items(), start=1):
        ws.cell(1, i, name).font = f(bold=True)
        for r, v in enumerate(values, start=2):
            ws.cell(r, i, v).font = f()
        refs[name] = f"Lists!${L(i)}$2:${L(i)}${len(values) + 1}"
    ws.sheet_state = "hidden"
    return refs


def draw_table(ws, groups, refs, example_rows, extra_refs=None):
    refs = dict(refs, **(extra_refs or {}))
    col = 1
    cols = []
    for gi, (gname, gcols) in enumerate(groups):
        start = col
        for c in gcols:
            cols.append((col, c))
            col += 1
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=col - 1)
        cell = ws.cell(1, start, gname)
        cell.font = f(bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor=GROUP_FILLS[gi % len(GROUP_FILLS)])
        cell.alignment = Alignment(horizontal="left", vertical="center")
        for cc in range(start, col):
            ws.cell(1, cc).fill = PatternFill("solid", fgColor=GROUP_FILLS[gi % len(GROUP_FILLS)])

    col_group = {}
    for gi, (gname, gcols) in enumerate(groups):
        for c in gcols:
            col_group[c["key"] + "|" + c["header"]] = GROUP_FILLS[gi % len(GROUP_FILLS)]

    for ci, c in cols:
        base = platform_of(c["header"]) or col_group[c["key"] + "|" + c["header"]]
        h = ws.cell(2, ci, c["header"] + (" *" if c["req"] else ""))
        h.font = f(bold=True, color="FFFFFF")
        h.fill = PatternFill("solid", fgColor=base)
        h.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        h.border = BORDER
        hint = c["help"] or ""
        if c["default"] not in (None, ""):
            hint = (hint + " " if hint else "") + f"(Blank = {c['default']})"
        if c["list"]:
            hint = (hint + " " if hint else "") + "[dropdown]"
        t = ws.cell(3, ci, hint)
        t.font = f(italic=True, color="595959", size=8)
        t.alignment = Alignment(wrap_text=True, vertical="top")
        t.fill = REQ_FILL if c["req"] else OPT_FILL
        t.border = BORDER
        ws.column_dimensions[L(ci)].width = c["width"]
        if c["list"]:
            dv = DataValidation(type="list", formula1=f"={refs[c['list']]}", allow_blank=True,
                                showErrorMessage=True, errorStyle="warning",
                                error="Pick a value from the list", errorTitle="Not in list")
            dv.add(f"{L(ci)}4:{L(ci)}{DATA_ROWS}")
            ws.add_data_validation(dv)
        body = PatternFill("solid", fgColor=tint(base, 0.82))
        for r in range(4, DATA_ROWS + 1):
            cell = ws.cell(r, ci)
            cell.font = f()
            cell.fill = body
            cell.border = BORDER

    for ri, ex in enumerate(example_rows, start=4):
        for ci, c in cols:
            v = ex.get(c["key"])
            cell = ws.cell(ri, ci, v)
            cell.fill = EX_FILL
            cell.font = f(italic=True, color="375623")

    # red warning: a row has a Style Code but a required cell is still empty
    first_letter = L(cols[0][0])
    for ci, c in cols[1:]:
        if c["req"]:
            rng = f"{L(ci)}4:{L(ci)}{DATA_ROWS}"
            ws.conditional_formatting.add(rng, FormulaRule(
                formula=[f'AND(${first_letter}4<>"",{L(ci)}4="")'],
                fill=PatternFill("solid", bgColor="FFC7CE", fgColor="FFC7CE"), stopIfTrue=False))
    ws.row_dimensions[2].height = 32
    ws.row_dimensions[3].height = 40
    ws.freeze_panes = "C4" if ws.title == "Products" else "C4"
    return cols


def build(path):
    wb = Workbook()
    guide = wb.active
    guide.title = "How to Fill"
    refs = write_lists(wb)

    prod = wb.create_sheet("Products", 1)
    pcols = draw_table(prod, S.PRODUCT_GROUPS, refs, [EXAMPLE_PRODUCT])
    style_col = L([ci for ci, c in pcols if c["key"] == "style_code"][0])
    sizes = wb.create_sheet("Sizes", 2)
    draw_table(sizes, S.SIZE_GROUPS, refs, EXAMPLE_SIZES,
               {"style_codes": f"Products!${style_col}$4:${style_col}${DATA_ROWS}"})

    # Settings -----------------------------------------------------------
    st = wb.create_sheet("Settings", 3)
    st["A1"] = "PLATFORM SETTINGS  (fill once - used for every product)"
    st["A1"].font = f(bold=True, size=12, color="1F3864")
    for i, ch in enumerate(["Setting"] + S.CHANNELS):
        c = st.cell(3, 1 + i, ch)
        c.font = f(bold=True, color="FFFFFF", size=11)
        c.fill = PatternFill("solid", fgColor=PLATFORM_COLOURS.get(ch, "1F3864"))
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center")
    st.row_dimensions[3].height = 22
    r = 4
    for label, vals in S.CHANNEL_SETTINGS:
        st.cell(r, 1, label).font = f(bold=True)
        st.cell(r, 1).border = BORDER
        st.cell(r, 1).fill = PatternFill("solid", fgColor="DDEBF7")
        for i, v in enumerate(vals):
            c = st.cell(r, 2 + i, v)
            c.font = f(color="0000FF"); c.border = BORDER
            c.fill = PatternFill("solid", fgColor=tint(PLATFORM_COLOURS[S.CHANNELS[i]], 0.88))
            c.alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    st.cell(4, 6, "Meesho brand/address copied from your current Meesho file (Silver Kraft / SS Lifestyle). Change if needed.").font = f(italic=True, color="C00000", size=9)
    r += 1
    st.cell(r, 1, "GENERAL").font = f(bold=True, size=11, color="1F3864"); r += 1
    for i, h in enumerate(["Setting", "Value", "Note"]):
        c = st.cell(r, 1 + i, h); c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD_FILL
    r += 1
    for label, val, note in S.GENERAL_SETTINGS:
        st.cell(r, 1, label).font = f(bold=True)
        c = st.cell(r, 2, val); c.font = f(color="0000FF"); c.fill = REQ_FILL; c.border = BORDER
        st.cell(r, 3, note).font = f(italic=True, color="595959", size=9)
        r += 1
    r += 1
    st.cell(r, 1, "SIZE TABLE").font = f(bold=True, size=11, color="1F3864"); r += 1
    for i, h in enumerate(["Size", "Meesho Kids Weight (kg)"]):
        c = st.cell(r, 1 + i, h); c.font = f(bold=True, color="FFFFFF"); c.fill = HEAD_FILL
    r += 1
    for s in vocab.SIZES:
        st.cell(r, 1, s).font = f()
        c = st.cell(r, 2, vocab.MEESHO_KIDS_WEIGHT[s]); c.font = f(color="0000FF"); c.fill = REQ_FILL
        r += 1
    st.column_dimensions["A"].width = 32
    for colL in "BCDE":
        st.column_dimensions[colL].width = 34
    st.column_dimensions["F"].width = 40

    # Guide --------------------------------------------------------------
    lines = [
        ("WISH KARO - MASTER LISTING SHEET", "title"),
        ("Fill this once. The tool creates the Meesho, Flipkart, Myntra and Amazon upload files from it.", "sub"),
        ("", None),
        ("STEP 1 - Products sheet", "h"),
        ("One row per COLOUR of a design (not per size).", None),
        ("Style Code = colour-level code (bxa245pkblu). Parent Code = design code shared by all its colours (bxa245).", None),
        ("Columns marked * (yellow note under the heading) must be filled. Others are optional.", None),
        ("A cell turns RED if you typed a Style Code but left a required (*) cell empty in that row.", None),
        ("Pink/purple, blue, pink and orange headings are only for Meesho, Flipkart, Myntra and Amazon.", None),
        ("Use the dropdowns. Do not type your own spellings - the tool converts them to each platform's wording.", None),
        ("Prices: fill Price + MRP for every platform you want. Leave a platform blank and that product is skipped there.", None),
        ("", None),
        ("STEP 2 - Sizes sheet", "h"),
        ("One row per size. Pick the Style Code from the dropdown so it matches the Products sheet.", None),
        ("SKU is created automatically (StyleCode_Size, e.g. bxa245pkblu_3-6_Months). Type one only if it is different.", None),
        ("Measurements: Chest and Front Length are needed for Myntra and Meesho.", None),
        ("Price Override: only if this size costs more or less than the product price (e.g. bigger sizes Rs 100 more).", None),
        ("", None),
        ("STEP 3 - Settings sheet", "h"),
        ("Brand, addresses, stock defaults etc. Fill once; change only if something changes.", None),
        ("", None),
        ("STEP 4 - Send this file", "h"),
        ("Save and send. You get back 4 files ready to upload, plus a report listing anything to fix.", None),
        ("", None),
        ("The green EXAMPLE rows show the format. They are ignored - you can leave them or delete them.", "note"),
        ("Images: Meesho rejects Google Drive links. Use Meesho's image uploader links in the 'Meesho Image' columns.", "note"),
    ]
    step_colours = iter(["2E75B6", "00897B", "B7791F", "D81B60"])
    colour = "1F3864"
    for i, (text, kind) in enumerate(lines, start=2):
        c = guide.cell(i, 2, text)
        c.alignment = Alignment(vertical="center", indent=1)
        if kind == "title":
            c.font = f(bold=True, size=18, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3864")
            guide.row_dimensions[i].height = 34
        elif kind == "sub":
            c.font = f(size=11, color="1F3864", italic=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
            guide.row_dimensions[i].height = 22
        elif kind == "h":
            colour = next(step_colours)
            c.font = f(bold=True, size=12, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=colour)
            guide.row_dimensions[i].height = 24
        elif kind == "note":
            c.font = f(bold=True, color="9C0006")
            c.fill = PatternFill("solid", fgColor="FFC7CE")
            guide.row_dimensions[i].height = 20
        elif text:
            c.font = f(size=10)
            c.fill = PatternFill("solid", fgColor=tint(colour, 0.88))
            guide.row_dimensions[i].height = 18
    guide.sheet_properties.tabColor = "1F3864"
    prod.sheet_properties.tabColor = "2E7D32"
    sizes.sheet_properties.tabColor = "00897B"
    st.sheet_properties.tabColor = "B7791F"
    guide.column_dimensions["A"].width = 3
    guide.column_dimensions["B"].width = 120
    guide.sheet_view.showGridLines = False

    wb.active = 1
    wb.save(path)
    return path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "Wishkaro_Master_Listing.xlsx"
    print("Saved", build(out))
